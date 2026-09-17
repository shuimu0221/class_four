# -*- coding: utf-8 -*-
"""Assemble extracted slide-frame images into viewable .pptx decks (one image per slide)."""
import glob
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

EMU_PER_INCH = 914400


def build_deck(image_dir, out_path):
    files = sorted(glob.glob(os.path.join(image_dir, "*.jpg")))
    assert files, f"no images found in {image_dir}"
    w, h = Image.open(files[0]).size
    aspect = w / h
    slide_w_in = 13.333
    slide_h_in = slide_w_in / aspect
    prs = Presentation()
    prs.slide_width = Emu(int(slide_w_in * EMU_PER_INCH))
    prs.slide_height = Emu(int(slide_h_in * EMU_PER_INCH))
    blank = prs.slide_layouts[6]
    for f in files:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(f, 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(out_path)
    print(f"Saved {out_path}  ({len(files)} slides, {slide_w_in:.2f}x{slide_h_in:.2f} in)")


base = r"C:\Users\ziang.xu\Documents\class_four\extracted_ppt"
build_deck(os.path.join(base, "armor_slides"),
           os.path.join(r"C:\Users\ziang.xu\Documents\class_four", "Hello_Armor_原视频PPT提取.pptx"))
build_deck(os.path.join(base, "pnp_slides"),
           os.path.join(r"C:\Users\ziang.xu\Documents\class_four", "2025装甲板位姿解算_原视频PPT提取.pptx"))
