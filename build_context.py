import os
import re
import sys

from docx import Document
from pptx import Presentation

R = r"C:/Users/ziang.xu/Documents/sp/class_four"
YOLO = os.path.join(R, "lecture4/yolo")
OUT = os.path.join(R, "output/script_parts/_context.md")

# ---------------- §一  PPT 全文 (39 pages) ----------------
# 优先读刚生成的版本（仓库里那份可能正被 PowerPoint 占用、尚未更新）
GEN = os.path.expandvars(r"%TEMP%\mat_gen").replace("\\", "/")
_cand = [os.path.join(GEN, "Lecture4_HelloArmor_装甲板位姿解算_YOLO版.pptx"),
         os.path.join(R, "Lecture4_HelloArmor_装甲板位姿解算_YOLO版.pptx"),
         os.path.join(R, "Lecture4_HelloArmor_装甲板位姿解算.pptx")]
PPTX_PATH = next(p for p in _cand if os.path.exists(p))
prs = Presentation(PPTX_PATH)
slides = list(prs.slides)
sys.stderr.write(f"PPTX source: {PPTX_PATH}\n")


def shape_lines(shape, depth=0):
    out = []
    pad = "  " * depth
    if shape.shape_type == 6 and hasattr(shape, "shapes"):
        for s in shape.shapes:
            out += shape_lines(s, depth + 1)
        return out
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t:
                out.append(pad + t)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            out.append(pad + " | ".join(c.text.strip() for c in row.cells))
    return out


ppt_sections = []
for i, sl in enumerate(slides, 1):
    body = []
    for sh in sl.shapes:
        body += shape_lines(sh)
    # drop the footer line and the bare page number
    body = [b for b in body if not b.startswith("Hello Armor · 装甲板位姿解算")]
    body = [b for b in body if not re.fullmatch(r"\d{2}", b.strip())]
    title = body[1] if len(body) > 1 else (body[0] if body else "")
    ppt_sections.append((i, title, body))

# ---------------- §二 教案 ----------------
doc = Document(os.path.join(R, "Lecture4_HelloArmor_教案.docx"))
doc_paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
doc_tables = []
for ti, t in enumerate(doc.tables):
    rows = []
    for row in t.rows:
        rows.append([c.text.strip().replace("\n", " / ") for c in row.cells])
    doc_tables.append(rows)

# ---------------- §三 原视频覆盖清单 ----------------
# 这份清单来自历史素材（两份原视频课），无法从产物推导，所以内嵌在这里。
#
# ⚠ 引用幻灯片一律用【标题】而不是页码。用页码会有一个致命问题：每次重建都要
#    把旧页码重映射成新页码，而重映射不是幂等的——重建两次页码就会漂两次。
#    标题不随页数变化，因此天然幂等。
SEC3 = """## 三、两份原视频课的「必须覆盖清单」（用户硬性要求：内容都要讲）

以下内容来自《2025装甲板位姿解算_原视频PPT提取》与《Hello_Armor_原视频PPT提取》两份原始课件/录课，
**讲稿必须全部覆盖**。括号里给出对应幻灯片的**标题**（不是页码——页码会随课件改版变化，
按标题在 `一、PPT 全文` 里检索即可）：

**动机与直观（P0/P1）**
- 3 号车 vs 4 号车："相比于 4 号车，3 号车离我们更远"——人眼能判断距离（「只有这 4 个 2D 点，够瞄准吗？」）
- 同一块装甲板三种角度："图片中的三块装甲板有什么不同？"——朝向差异（同上页）
- 欧拉角按特定顺序定义：先偏航（绕上下轴）、再俯仰（绕左右轴）、后横滚（绕前后轴），每一步相对前一步的结果（「位姿 = 位置 + 朝向」）
- "看出"三维信息 = 距离远近 + 朝向左右 → 更准确地说：位置用三维向量、旋转用欧拉角 → "程序能解算这些信息吗？"（「位姿 = 位置 + 朝向」「先认清两个坐标系」）

**PnP 与 API（P2）**
- PnP：n 组一一对应的点 + 相机内参/畸变 → 解 R、t（「Perspective-n-Points 问题」）
- solvePnP 逐参数讲解（「cv::solvePnP 函数签名精讲」）
- tvec = 装甲板坐标系原点在相机系下的位置；**rvec 留悬念："先写代码，留个悬念，后面再来解答"**（「输出到底是什么？」）

**rvec 揭秘（P4）——核心章节，必须讲透**
- "转一转，看一看"：观察装甲板转动时 rvec 方向和大小怎么变（「转一转，看一看」）
  ⚠ 原课是在**终端**滚动打印、且数值量级是 `roll: 2.01 → 6.40`（疑为角度制的 roll）。
  **本届程序不往终端打任何东西**，tvec/rvec/欧拉角都画在**画面左上角**、原地刷新，且**打的是弧度**。
  讲稿**不得**引用 2.01/6.40 这类数值，也不得说"终端在滚动"。
- rvec 定义：方向 = 旋转轴，模长 = 转角（弧度）；3 个数表示任意旋转（「rvec 到底是什么」）
- 概念图：旋转向量 rvec --罗德里格斯公式/cv::Rodrigues（可逆）--> 旋转矩阵 rmat --反三角函数--> 欧拉角；rmat --> 四元数（「旋转的几种表示法，一张图看懂」）
- **INT_YXZ 欧拉角公式：θ₁ = arctan2(m₁₃, m₃₃)，θ₂ = −arcsin(m₂₃)，θ₃ = arctan2(m₂₁, m₂₂)**（「从旋转矩阵提取欧拉角（INT_YXZ 约定）」）
- 关键对应关系：公式里的 m₁₃ 是 rmat 第 1 行第 3 列，但 **cv::Mat 下标从 0 开始**，所以代码里写 `rmat.at<double>(0, 2)`——这是学生最容易错的地方（同上页）
- `std::atan2(double, double)` 是 C++ 标准库函数（原课专门给过这个提示）（同上页）
- 弧度换算：**本届程序打弧度**；57.3 (≈180/π) 只作为"想看度数就这么换算"的补充（同上页注释）
- 万向锁：欧拉角表示不唯一、特定姿态丢一个自由度；四元数 (w,x,y,z) 无奇异点、插值稳健、不直观（「欧拉角的坑：万向锁（Gimbal Lock）」）
- **quaternions.online 现场交互演示**：Quaternion 的 W/X/Y/Z 输入框 + Euler Angles 的 X/Y/Z 和 XYZ-Order 下拉框，红 x / 绿 y / 蓝 z 的 3D 网格实时联动——原 Hello Armor 课现场演示过浏览器这个网站（同上页，讲稿应作为现场/课后演示展开）

**动手实验（P3/P5）**
- 点序约定：左上、右上、右下、左下（自左上顺时针）；两个点集顺序必须一致（「实验环境 & 点位约定」）
- **点序陷阱**：`tasks/armor.hpp` 里 `Armor::points` 的注释是**错的**（写的是左上、左下、右下、右上）。
  本届新增一整页专讲这件事（「点序陷阱：注释是错的」），并配一个自查工具「重投影误差」（「重投影误差：给点序配一把尺」）
- Task 01~03 填空思路与提示，不给答案（「Task 01 · 填写 object_points」至「Task 03 · 调用 cv::solvePnP」）
- 经典报错：本届程序的第一个报错**不是视频、是配置**——`YAML::BadFile` / `what(): bad file: configs/yolo.yaml`，
  在 build/ 目录里直接 ./main 导致；`cd` 回 `lecture4/yolo` 再 `./build/main`（「现场演示 & Debug 小贴士」）
- Task 04：`tvec.at<double>(0)` 取出 double；fmt::format 与 draw_text（「Task 04 · 把 tvec / rvec 打印到画面上」）
- Task 05：Rodrigues + 反三角函数 + 显示（「Task 05 · rvec → rmat → 欧拉角」）
- 观察 yaw/pitch/roll 实时变化；邀请学生上台；提问"前后移动（不转动）装甲板，tvec 和 rvec 分别怎么变？"
  ——答：tvec 变、rvec 基本不变（「现场演示 & 互动」）

**坐标系全景（P6）**
- 相机坐标系：与相机刚性连接、随相机平移旋转、原点 = 镜头光心；"如果相机安装有倾斜，并且和枪管不在同一个位置，会怎么样？"（「我真的对准了吗？」）
- "实际车的运动状态千奇百怪，我们需要装甲板姿态来推算出它的旋转中心，进而对运动进行拟合"（「工程上为什么必须要"姿态"，不只是"位置"」）
- 手眼标定：获取机器人本体坐标系→相机坐标系的变换（机械臂 AX=XB 标定环）（「完整的坐标变换链条」）
- IMU："机器人本体控制依赖于 imu（陀螺仪），需要知道 imu 到本体坐标系之间的变换"（同上页）
- 完整链条：像素坐标系 → 相机坐标系 → 机器人本体坐标系 → IMU 坐标系，分别靠 solvePnP / 手眼标定 / 实际安装（同上页）

**收尾（P7）**
- 8 条回顾（「本讲回顾」）；三档作业（「课后作业」）；Thanks + 下一讲预告（「Thanks · Q&A」）
- 原课收尾还带过一遍课程大纲（可作 P7 口头回顾素材）。注意**本届的编号不同**：
  Lecture 2 = Hello C++&&OOP、Lecture 3 = Hello Modern C++、Lecture 4 = 本讲、Lecture 5 = Hello Kalman && Target。
  不要说"上一讲是 OOP"——OOP 是上上讲。

**可引用的课后参考链接（原课件出现过，可作"放群里"的口头补充）**
- 万向锁知乎文章：https://zhuanlan.zhihu.com/p/9135205633
- 线性代数（旋转矩阵背景）B 站视频：https://www.bilibili.com/video/BV1ns41167b9
- OpenCV 四元数类文档 `cv::Quat`：https://docs.opencv.org/4.x/d4/d4a/classcv_1_1Quat.html
- 交互演示网站：quaternions.online

"""

# ---------------- §四 工程代码 ----------------
def read(rel):
    p = os.path.join(YOLO, rel)
    return open(p, encoding="utf-8").read() if os.path.exists(p) else f"<missing: {rel}>"


main_src = read("src/main.cpp")
cmake = read("CMakeLists.txt")
io_cam = read("io/camera.hpp")
yamlcfg = read("configs/yolo.yaml")

# ---------------- assemble ----------------
L = []
w = L.append

w("# 讲稿撰写 · 上下文包（唯一事实来源）v4 —— YOLO 版")
w("")
w("本文件是撰写《Lecture 4 Hello Armor 装甲板位姿解算》逐字讲稿的唯一素材来源。")
w("")
w("> ⚠️ **基准材料**（以本文件为准，不得引用其他版本）：")
w("> - **课件**：根目录《Lecture4_HelloArmor_装甲板位姿解算_YOLO版.pptx》（**39 页**）")
w("> - **教案**：根目录《Lecture4_HelloArmor_教案.docx》")
w("> - **工程代码**：`lecture4/yolo`（学生版 `src/main.cpp` + 参考版 `answer/main.cpp`，同一个 CMake 工程）")
w("> - **配套文档**：`lecture4/yolo/docs/keypoint_order.md`（点序陷阱的完整证据链）、`lecture4/yolo/README.md`")
w("> - **内容覆盖要求**：两份原视频 PPT《2025装甲板位姿解算_原视频PPT提取》《Hello_Armor_原视频PPT提取》的内容**全部都要讲**（含欧拉角、万向锁、四元数——这些是本讲核心教学内容，不是了解内容）")
w("> - **本届改用 YOLO 版程序**：旧的手写传统 CV 版 `lecture4/class/` **不再讲授**，只作为代码保留。")
w(">   `lecture4/yolo/tasks/` 与 `tools/` 是从学生 lecture2 作业**逐字复制**过来的，一行没改。")
w("> - `output/` 目录下带 `_v4/_v5/_v2` 后缀的旧修订版 PPT/教案**已废弃**，与当前代码不符，不得引用。")
w("")
w("---")
w("")

# §〇
w("## 〇、任务白名单（最高优先级约束）")
w("")
w("**本讲只有 5 个 Task，全部位于同一个文件 `lecture4/yolo/src/main.cpp`（学生版）。**")
w("（`answer/main.cpp` 是教师参考答案，与学生版逐 Task 对应；它**默认不编译**，教师需 `-DBUILD_ANSWER=ON`。）")
w("")
w("| Task | 学生要做什么 |")
w("|---|---|")
w("| **Task 01** | 填写 `object_points`：装甲板局部坐标系下 4 个点的 3D 坐标（模板被注释着，要先解注释再填）。点序：左上、右上、右下、左下 |")
w("| **Task 02** | 填写 `img_points`：4 个 2D 像素点，**取自 `armor.points.at(0)`~`at(3)`** |")
w("| **Task 03** | 调用 `cv::solvePnP(object_points, img_points, camera_matrix, distort_coeffs, rvec, tvec);` |")
w("| **Task 04** | 改写 `draw_text` 参数，把解得的 tvec / rvec 打印到画面上（现在只打印 0.0） |")
w("| **Task 05** | 用 `cv::Rodrigues(rvec, rmat)` 转旋转矩阵，再用反三角函数从 rmat 提取欧拉角 yaw/pitch/roll 并显示 |")
w("")
w("**工程里没有的东西（讲稿不得布置、不得要求实现）：**")
w("- ❌ 距离 / 瞄准角 `distance_m` / `aim_angles_deg`")
w("- ❌ `PoseSolver` 类、`pose_solver_test`、ctest 验收")
w("- ❌ ROI、动态 ROI、颜色二次估计（`configs/yolo.yaml` 里的 `threshold` 与 `use_traditional` 是历史遗留字段，代码已不用）")
w("- ❌ 学生要自己实现检测器（`tasks/` 是发的成品，与 lecture2 作业里那份一致）")
w("")
w("**本届新增、必须讲到的两页：**")
w("- **第 18 页 · 点序陷阱**：`tasks/armor.hpp` 里 `Armor::points` 的注释是错的")
w("- **第 31 页 · 重投影误差自查**：`tools/pnp_check.hpp`，判读口径「个位数像素=对，两位数=错」")
w("")
w("---")
w("")

# §一
w(f"## 一、PPT 全文（共 {len(slides)} 页，逐页）")
w("")
for i, title, body in ppt_sections:
    w(f"### 第 {i} 页 —— {title}")
    w("")
    for b in body:
        w(b)
    w("")
w("---")
w("")

# §二
w("## 二、教案（教学设计与时间分配）")
w("")
w("### 基本信息")
w("")
for row in doc_tables[0]:
    w(f"- {row[0]}：{row[1]}")
    w(f"- {row[2]}：{row[3]}")
w("")
w("### 教学过程与时间分配（共 90 分钟；页码为 39 页课件实际页序）")
w("")
for row in doc_tables[1]:
    w("| " + " | ".join(row) + " |")
w("")
w("### 教学目标、重点与难点")
w("")
grab = False
for p in doc_paras:
    if p.startswith("一、教学目标"):
        grab = True
    if p.startswith("二、教学过程"):
        break
    if grab:
        w(p)
w("")
w("### Task 01~05 参考答案（教师用，**不对学生展示**；讲稿只讲思路与提示）")
w("")
grab = False
for p in doc_paras:
    if p.startswith("三、Task"):
        grab = True
        continue
    if p.startswith("四、演示与教具清单"):
        break
    if grab:
        w(p)
w("")
w("### 演示与教具")
w("")
grab = False
for p in doc_paras:
    if p.startswith("四、演示与教具清单"):
        grab = True
        continue
    if p.startswith("五、作业布置"):
        break
    if grab:
        w(p)
w("")
w("### 作业布置（三档）")
w("")
grab = False
for p in doc_paras:
    if p.startswith("五、作业布置"):
        grab = True
        continue
    if p.startswith("六、教学反思"):
        break
    if grab:
        w(p)
w("")
w("---")
w("")

# §三
w(SEC3.rstrip())
w("")
w("---")
w("")

# §四
w("## 四、工程代码（讲稿涉及代码必须与这里一致）")
w("")
w("### 4.1 目录与构建")
w("")
w("```")
w("lecture4/yolo/")
w("├── CMakeLists.txt      # 顶层；依赖 OpenCV/fmt/Eigen3/yaml-cpp/spdlog/OpenVINO")
w("├── configs/yolo.yaml   # 检测与取帧配置")
w("├── assets/             # yolov5.xml/.bin + video.avi（约 60MB）")
w("├── docs/               # keypoint_order.md + verify_keypoint_order.py")
w("├── tools/  tasks/      # 逐字复制自 lecture2 作业，未改动")
w("├── io/                 # 取帧：按 config 在 video / camera 间切换")
w("├── src/main.cpp        # 学生填空（Task 01~05）")
w("└── answer/main.cpp     # 参考实现（默认不编译）")
w("```")
w("")
w("```bash")
w("cd lecture4/yolo")
w("cmake -B build                 # 教师演示另加 -DBUILD_ANSWER=ON")
w("cmake --build build -j")
w("./build/main                  # 必须在 lecture4/yolo/ 下运行")
w("```")
w("")
w("顶层 CMakeLists.txt 要点：")
w("")
w("```cmake")
w(cmake.strip())
w("```")
w("")
w("### 4.2 `src/main.cpp`（学生版全文）")
w("")
w("```cpp")
w(main_src.strip())
w("```")
w("")
w("### 4.3 `io/camera.hpp`（取帧接口）")
w("")
w("```cpp")
w(io_cam.strip())
w("```")
w("")
w("### 4.4 `configs/yolo.yaml`")
w("")
w("```yaml")
w(yamlcfg.strip())
w("```")
w("")
w("### 4.5 检测器与关键点顺序")
w("")
w("`tasks/yolos/yolov5.cpp` 走 OpenVINO 推理 `assets/yolov5.xml`，输出 `[1, 25200, 22]`：")
w("`0–7` 四个关键点、`8` objectness、`9–12` 颜色独热、`13–21` 类别独热。")
w("`parse()` 用 4 个关键点的外接矩形当 box，并按 `[0,3,2,1]` 置换后交给 `Armor` 构造函数。")
w("")
w("**最终 `Armor.points` 的顺序是 左上、右上、右下、左下**（自左上顺时针）。")
w("`tasks/armor.hpp` 里的注释写的是「左上、左下、右下、右上」——**注释是错的**，原样保留（与学生机器上一致）。")
w("证据链见 `lecture4/yolo/docs/keypoint_order.md`；实测重投影误差中位约 2.7 px。")
w("")
w("**`armor.left` / `armor.right` 在本讲是空的**：YOLO 构造函数只初始化 `confidence` / `box` / `points`，")
w("`left`/`right` 走 `Lightbar() {}` 默认构造，四个 `cv::Point2f` 全是 `(0,0)`。")
w("照旧材料写 `armor.left.top` 会拿到四个 `(0,0)`，solvePnP 失败或返回垃圾，**且不报错**。")
w("")
w("### 4.6 跑通后的画面")
w("")
w("学生版画三行（tvec / rvec / euler angles，未填空时全是 0.0）；参考版多一行 `reproj err`。")
w("程序**不往终端打任何数值**。取帧默认读 `assets/video.avi`（`loop: true`，25 秒一轮，内容是手持数字「2」的装甲板）。")
w("按 `q` 退出，按**空格**暂停。")

open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")
sys.stderr.write(f"wrote {len(L)} lines\n")
sys.stderr.write(f"ppt pages: {len(ppt_sections)}  docx paras: {len(doc_paras)}  sec3: {len(SEC3)} chars\n")
